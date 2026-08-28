#!/usr/bin/env python3
"""
Generate a clean, plain, unformatted Excel report of June and July unrenewed licenses.
No custom fills, no merged banner rows, no background colors.
"""

import json
import openpyxl

json_file = "/Users/lambert/Desktop/fast-api/ewura-migration/expired_data.json"
with open(json_file, "r") as f:
    data = json.load(f)

june_data = data.get("june", [])
july_data = data.get("july", [])

wb = openpyxl.Workbook()
wb.remove(wb.active)

headers = [
    "Application Number",
    "Approval Number (License No)",
    "Company / Facility / Applicant Name",
    "Sector",
    "License Type Code",
    "Category Name",
    "Category Type",
    "Region",
    "District",
    "Ward",
    "Zone",
    "Effective Date",
    "Expire Date",
    "Status As Of Month-End",
    "Reminder Window (Days)",
    "Intimate Date (Reminder Start)",
    "Origin"
]

def populate_sheet_simple(ws, title, rows):
    ws.title = title
    ws.views.sheetView[0].showGridLines = True
    
    # Row 1: Headers
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    
    # Row 2+: Data
    for row_idx, r in enumerate(rows, start=2):
        row_vals = [
            r.get("application_number") or "",
            r.get("approval_no") or "",
            r.get("company_or_facility_name") or "",
            r.get("sector") or "",
            r.get("license_type") or "",
            r.get("category_name") or "",
            r.get("category_type") or "",
            r.get("region") or "",
            r.get("district") or "",
            r.get("ward") or "",
            r.get("zone_name") or "",
            r.get("effective_date") or "",
            r.get("expire_date") or "",
            r.get("renewal_status") or "",
            r.get("reminder_window_days"),
            r.get("intimate_date") or "",
            r.get("origin") or ""
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

ws_june = wb.create_sheet("June 2026 Report")
populate_sheet_simple(ws_june, "June 2026 Report", june_data)

ws_july = wb.create_sheet("July 2026 Report")
populate_sheet_simple(ws_july, "July 2026 Report", july_data)

out_file = "/Users/lambert/Desktop/fast-api/ewura-migration/Expired_Licenses_June_July_2026.xlsx"
wb.save(out_file)
print(f"Successfully saved plain unformatted Excel report at {out_file}")
